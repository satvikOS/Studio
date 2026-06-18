#!/usr/bin/env python
"""Assemble render/*.png into a 16:9 PPTX (full-bleed), embedding the demo video."""
import glob, os
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

BD = os.path.expanduser('~/archdisc-deck')
EMU_W, EMU_H = 12192000, 6858000      # 13.333in x 7.5in (16:9)
PXEMU = EMU_W // 1280                  # 9525 EMU per logical px (1280x720 authoring grid)

prs = Presentation()
prs.slide_width = Emu(EMU_W); prs.slide_height = Emu(EMU_H)
blank = prs.slide_layouts[6]

# presenter talk track, embedded as speaker notes (one tight paragraph each)
SCRIPTS = {
 '01': "We're ArchDisc, and our whole idea is four words. Describe it, Archie builds it. We have two apps. "
       "Studio for 3D art, and Forge for real engineering. And one AI, Archie, that builds in both. I'm Satvik, "
       "this is Jeff. Give us a few minutes, and we'll show you how anyone can make real things, just by saying what they want.",
 '02': "Here's what bugs us. Making real things is still painfully slow. An expert can lose hours on something as small "
       "as a bracket or a coffee mug. And if you never trained for years, you can't even start. So most ideas just die "
       "in someone's head. That is the gap we close.",
 '03': "So how does it work. Two apps, one brain. Studio gives you the power of Blender, Maya and Houdini. Forge gives you "
       "a real engineering kernel, for parts you can actually build. And Archie runs through both. You don't learn the tool. "
       "You say what you want, and it builds it.",
 '04': "Don't take our word for it. Watch. One sentence goes in. Archie plans it, builds it, and brings it to life, right on "
       "the machine. No cloud. No waiting. This is real, and it runs today.",
 '05': "Why us, and why now. Four things. It is a real kernel, so you get parts you can build, not just a pretty picture. "
       "It is yours and private, free on your own machine. It is one model across art and engineering, and no one else spans both. "
       "And it gets smarter every time someone uses it. All of that, in a 3D software industry already worth thirty billion a year.",
 '06': "And this is not a thin wrapper on someone else's AI. There is real engineering underneath. A native kernel for Forge. "
       "A full 3D engine for Studio. And our own models, trained on millions of our own samples, running right on your laptop. "
       "We built the hard parts ourselves.",
 '07': "How we make money is simple, and you have seen it work before. The same as GitHub and Vercel. Free for anyone creating "
       "on their own. They build anything, the source stays ours. When their work outgrows one machine, they pay for Pro, for "
       "cloud power and storage. And teams pay the most, for working together, security, and their own custom models.",
 '08': "Our path is bottom up. We land with students, makers and indie engineers, the people who love to try things for free. "
       "Competitions like ClawComp bring us the first thousands. Those people pull us into their studios, and studios pull in "
       "the big companies. And every user makes Archie better for the next one. It compounds.",
 '09': "So here is what we believe. If you can dream it, and you dare to create it, you can make worlds. That used to take "
       "years. Now it takes a sentence. This is the beginning of what we call vibe designing. We would love for you to build it "
       "with us. Thank you.",
}

pngs = sorted(glob.glob(os.path.join(BD, 'render', '0[0-9]-*.png')))
assert len(pngs) == 9, f'expected 9 slides, got {len(pngs)}: {[os.path.basename(p) for p in pngs]}'

# demo video frame rect in logical px (must match .vframe in 04-demo.html)
VF = dict(left=260, top=190, w=760, h=428)
VIDEO = os.path.join(BD, 'assets', 'demo.mp4')

for p in pngs:
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(p, Emu(0), Emu(0), width=Emu(EMU_W), height=Emu(EMU_H))
    note = SCRIPTS.get(os.path.basename(p)[:2])
    if note:
        s.notes_slide.notes_text_frame.text = note
    if os.path.basename(p).startswith('04-demo') and os.path.exists(VIDEO):
        # poster = crop of the rendered slide at the frame rect (render is 2x)
        poster = os.path.join(BD, 'render', '_demo_poster.png')
        im = Image.open(p); sx = im.width / 1280.0
        box = (int(VF['left']*sx), int(VF['top']*sx),
               int((VF['left']+VF['w'])*sx), int((VF['top']+VF['h'])*sx))
        im.crop(box).save(poster)
        s.shapes.add_movie(VIDEO,
            Emu(VF['left']*PXEMU), Emu(VF['top']*PXEMU),
            Emu(VF['w']*PXEMU), Emu(VF['h']*PXEMU),
            poster_frame_image=poster, mime_type='video/mp4')

cp = prs.core_properties
cp.title = 'ArchDisc — Investor Pitch'
cp.author = 'Satvik Adyanthaya · Jeff Munkondaya'
cp.comments = 'ClawComp x Link Ventures'

out = os.path.join(BD, 'out', 'ArchDisc-LinkVentures.pptx')
prs.save(out)
print('saved', out, '·', len(pngs), 'slides ·', os.path.getsize(out)//1024, 'KB · video embedded')
